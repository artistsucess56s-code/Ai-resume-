"""ATS Resume Checker - Streamlit + Gemini Flash.

Upload a resume (PDF / DOCX / TXT), optionally paste a job description,
and get an ATS score plus concrete suggestions for improvement.
"""

import io
import os
import re
from typing import List, Optional

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pydantic import BaseModel, Field
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-3.5-flash"
MAX_RESUME_CHARS = 30_000
MAX_JD_CHARS = 10_000


# --------------------------------------------------------------------------
# Structured output schema for Gemini
# --------------------------------------------------------------------------
class Improvement(BaseModel):
    section: str = Field(description="Resume section this applies to, e.g. Summary, Experience")
    issue: str = Field(description="What is wrong or weak")
    suggestion: str = Field(description="Specific, actionable fix")
    example: Optional[str] = Field(default=None, description="Optional rewritten example line")
    priority: str = Field(description="High, Medium or Low")


class ATSAnalysis(BaseModel):
    keyword_score: int = Field(description="0-100: relevance/keyword coverage for the role")
    content_score: int = Field(description="0-100: quality of achievements, metrics, clarity")
    structure_score: int = Field(description="0-100: ATS-friendly structure and readability")
    summary: str = Field(description="2-3 sentence overall assessment")
    strengths: List[str]
    matched_keywords: List[str]
    missing_keywords: List[str]
    improvements: List[Improvement]


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------
def extract_text(uploaded_file) -> str:
    """Extract plain text from an uploaded PDF, DOCX or TXT file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# --------------------------------------------------------------------------
# Rule-based ATS checks (fast, deterministic, no AI)
# --------------------------------------------------------------------------
SECTION_PATTERNS = {
    "Experience": r"\b(work experience|professional experience|experience|employment history)\b",
    "Education": r"\b(education|academic background|qualifications)\b",
    "Skills": r"\b(skills|technical skills|core competencies|technologies)\b",
    "Summary": r"\b(summary|profile|objective|about me)\b",
    "Projects": r"\b(projects|personal projects|key projects)\b",
}

ACTION_VERBS = {
    "led", "built", "developed", "designed", "created", "implemented", "managed",
    "improved", "increased", "reduced", "delivered", "launched", "optimized",
    "automated", "analyzed", "achieved", "established", "streamlined", "drove",
    "architected", "mentored", "coordinated", "deployed", "engineered", "owned",
}


def rule_based_checks(text: str) -> dict:
    """Return a format score (0-100) and a list of pass/fail checks."""
    lower = text.lower()
    words = re.findall(r"\b\w+\b", text)
    word_count = len(words)
    checks = []  # (label, passed, weight, detail)

    has_email = bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text))
    checks.append(("Email address present", has_email, 10, "Recruiters and ATS need a contact email."))

    has_phone = bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text))
    checks.append(("Phone number present", has_phone, 5, "Add a phone number."))

    has_link = bool(re.search(r"(linkedin\.com|github\.com|portfolio|https?://)", lower))
    checks.append(("LinkedIn / GitHub / portfolio link", has_link, 5, "Add a professional profile link."))

    for section, pattern in SECTION_PATTERNS.items():
        weight = 12 if section in ("Experience", "Education", "Skills") else 4
        found = bool(re.search(pattern, lower))
        checks.append((f"'{section}' section detected", found, weight, f"Add a clearly labelled {section} heading."))

    good_length = 300 <= word_count <= 1000
    checks.append((f"Length is reasonable ({word_count} words)", good_length, 10,
                   "Aim for roughly 300-1000 words (1-2 pages)."))

    bullet_lines = [l for l in text.splitlines() if re.match(r"^\s*[-•*▪●◦‣·]\s*\S", l)]
    checks.append(("Uses bullet points", len(bullet_lines) >= 5, 8,
                   "Use bullet points to describe achievements."))

    verb_hits = sum(1 for w in words if w.lower() in ACTION_VERBS)
    checks.append(("Uses strong action verbs", verb_hits >= 5, 8,
                   "Start bullets with verbs like Led, Built, Improved."))

    has_numbers = len(re.findall(r"\b\d+(?:[.,]\d+)?\s?(%|k|m|x|\+)?", text)) >= 5
    metric_hits = len(re.findall(r"\d+(?:[.,]\d+)?\s?(?:%|k\b|m\b|x\b|\+)", lower))
    checks.append(("Quantified achievements (numbers / %)", has_numbers and metric_hits >= 2, 8,
                   "Add metrics, e.g. 'reduced load time by 35%'."))

    total_weight = sum(c[2] for c in checks)
    earned = sum(c[2] for c in checks if c[1])
    score = round(100 * earned / total_weight) if total_weight else 0
    return {"score": score, "checks": checks, "word_count": word_count}


# --------------------------------------------------------------------------
# Gemini call
# --------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n\"\"\"\n{job_description[:MAX_JD_CHARS]}\n\"\"\"\n"
        "Score keyword_score by how well the resume matches this job description. "
        "missing_keywords must be important terms from the job description absent from the resume."
        if job_description.strip()
        else "No job description was provided. Judge keyword_score on general industry-standard "
        "keywords for the role this resume targets, and list missing_keywords that are commonly expected."
    )
    return f"""You are an expert technical recruiter and ATS (Applicant Tracking System) specialist.
Evaluate the resume below. Be honest and specific; do not inflate scores.
Scores are integers from 0 to 100.
Return 5 to 8 improvements, ordered by priority (High first). Where helpful, include a rewritten
example line, but NEVER invent employers, degrees, dates or metrics the candidate did not state;
use placeholders like [X%] when a number is needed.

{jd_block}

RESUME:
\"\"\"
{resume_text[:MAX_RESUME_CHARS]}
\"\"\"
"""


def analyze_with_gemini(resume_text: str, job_description: str, api_key: str, model: str) -> ATSAnalysis:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=ATSAnalysis,
            temperature=0.2,
        ),
    )
    parsed = getattr(response, "parsed", None)
    if isinstance(parsed, ATSAnalysis):
        return parsed
    if not response.text:
        raise RuntimeError("Gemini returned an empty response (it may have been blocked). Try again.")
    return ATSAnalysis.model_validate_json(response.text)


def clamp(value: int) -> int:
    return max(0, min(100, int(value)))


def get_api_key() -> str:
    try:
        secret = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:  # no secrets file locally
        secret = ""
    return secret or os.environ.get("GEMINI_API_KEY", "")


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------
def score_label(score: int) -> str:
    if score >= 80:
        return "🟢 Excellent"
    if score >= 60:
        return "🟡 Good, needs polish"
    return "🔴 Needs work"


def render_results(rules: dict, ai: ATSAnalysis) -> None:
    keyword, content, structure = clamp(ai.keyword_score), clamp(ai.content_score), clamp(ai.structure_score)
    ai_score = 0.4 * keyword + 0.35 * content + 0.25 * structure
    final = clamp(round(0.7 * ai_score + 0.3 * rules["score"]))

    st.header(f"ATS Score: {final}/100")
    st.progress(final / 100, text=score_label(final))
    st.write(ai.summary)

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Keywords", f"{keyword}/100")
    c2.metric("Content", f"{content}/100")
    c3.metric("Structure (AI)", f"{structure}/100")
    c4.metric("Format checks", f"{rules['score']}/100")

    tab1, tab2, tab3, tab4 = st.tabs(["🛠️ Improvements", "🔑 Keywords", "✅ Format checks", "💪 Strengths"])

    with tab1:
        order = {"high": 0, "medium": 1, "low": 2}
        for imp in sorted(ai.improvements, key=lambda i: order.get(i.priority.lower(), 3)):
            icon = {"high": "🔴", "medium": "🟠", "low": "🟡"}.get(imp.priority.lower(), "⚪")
            with st.expander(f"{icon} {imp.section} - {imp.issue}", expanded=imp.priority.lower() == "high"):
                st.markdown(f"**Fix:** {imp.suggestion}")
                if imp.example:
                    st.markdown("**Example:**")
                    st.code(imp.example, language=None)

    with tab2:
        k1, k2 = st.columns(2)
        with k1:
            st.subheader("Matched")
            st.write(", ".join(f"`{k}`" for k in ai.matched_keywords) or "None found")
        with k2:
            st.subheader("Missing")
            st.write(", ".join(f"`{k}`" for k in ai.missing_keywords) or "None 🎉")

    with tab3:
        st.caption(f"Word count: {rules['word_count']}")
        for label, passed, _w, detail in rules["checks"]:
            if passed:
                st.success(label)
            else:
                st.warning(f"{label} - {detail}")

    with tab4:
        for s in ai.strengths:
            st.markdown(f"- {s}")


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume, get an ATS score and specific ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if api_key:
            st.success("API key loaded from secrets/environment.")
        else:
            api_key = st.text_input("Gemini API key", type="password",
                                    help="Get a free key at https://aistudio.google.com/apikey")
        model = st.text_input("Gemini model", value=os.environ.get("GEMINI_MODEL", DEFAULT_MODEL))
        st.info("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")

    col_left, col_right = st.columns(2)
    with col_left:
        uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    with col_right:
        job_description = st.text_area("Job description (optional, improves keyword matching)", height=180)

    if not st.button("Analyze resume", type="primary", disabled=uploaded is None):
        return

    if not api_key:
        st.error("Please enter your Gemini API key in the sidebar.")
        return

    try:
        with st.spinner("Reading resume..."):
            text = extract_text(uploaded)
    except Exception as exc:
        st.error(f"Could not read the file: {exc}")
        return

    if len(text) < 100:
        st.error("Very little text was extracted. If your resume is a scanned image, "
                 "export it as a text-based PDF or DOCX and try again.")
        return

    rules = rule_based_checks(text)

    try:
        with st.spinner("Analyzing with Gemini..."):
            analysis = analyze_with_gemini(text, job_description, api_key, model.strip() or DEFAULT_MODEL)
    except Exception as exc:
        st.error(f"Gemini analysis failed: {exc}")
        return

    render_results(rules, analysis)


if __name__ == "__main__":
    main()
